"""
Contact extraction engine.
Supports .xlsx/.xls/.csv, .docx, .pdf, .txt.
Handles Forbes Marshall Google Contacts export format and arbitrary variants.
"""
from __future__ import annotations

import io
import re
import json
import logging
import os
from typing import Any

import pandas as pd

log = logging.getLogger(__name__)

# ── Column aliases ────────────────────────────────────────────────────────────
EMAIL_ALIASES = [
    "email", "email address", "e-mail", "e-mail 1 - value", "e-mail 2 - value",
    "mail id", "email id", "emailaddress", "mail address",
]
FIRST_NAME_ALIASES = ["first name", "firstname", "given name"]
LAST_NAME_ALIASES = ["last name", "lastname", "surname", "family name"]
NAME_ALIASES = [
    "name", "full name", "employee name", "contact name",
    "member name", "member",
] + FIRST_NAME_ALIASES + LAST_NAME_ALIASES

BRANCH_ALIASES = ["branch", "branch name", "location", "office"]
SALESREP_ALIASES = [
    "sales rep", "sales representative", "salesperson", "sales person", "rep",
    "account manager",
]
DEPARTMENT_ALIASES = ["department", "organization department", "dept"]
COMPANY_ALIASES = ["organization name", "company", "company name", "organization"]
TITLE_ALIASES = ["organization title", "title", "job title", "position", "role"]

_EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")


def _norm(col: str) -> str:
    return str(col).strip().lower()


def _find_col(columns: list[str], aliases: list[str]) -> str | None:
    normed = {_norm(c): c for c in columns}
    for alias in aliases:
        if alias in normed:
            return normed[alias]
    return None


def _extract_from_dataframe(df: pd.DataFrame, source_file: str = "") -> list[dict]:
    """Extract contacts from a DataFrame with flexible column detection."""
    cols = list(df.columns)
    normed_cols = [_norm(c) for c in cols]

    # Detect email column
    email_col = _find_col(cols, EMAIL_ALIASES)
    # Also check secondary email
    email2_col = _find_col(cols, ["e-mail 2 - value", "email2", "secondary email"])

    # Detect name — Forbes Marshall format has First Name + Last Name
    first_col = _find_col(cols, FIRST_NAME_ALIASES)
    last_col = _find_col(cols, LAST_NAME_ALIASES)
    name_col = _find_col(cols, ["name", "full name", "employee name", "contact name",
                                  "member name", "member"])

    branch_col = _find_col(cols, BRANCH_ALIASES)
    salesrep_col = _find_col(cols, SALESREP_ALIASES)
    dept_col = _find_col(cols, DEPARTMENT_ALIASES)
    company_col = _find_col(cols, COMPANY_ALIASES)
    title_col = _find_col(cols, TITLE_ALIASES)

    if not email_col:
        log.warning("No email column found in contact sheet. Columns: %s", cols[:10])
        return []

    results = []
    for _, row in df.iterrows():
        email = str(row.get(email_col, "") or "").strip()
        if not email or email.lower() in ("nan", "none", ""):
            # Try secondary email
            if email2_col:
                email = str(row.get(email2_col, "") or "").strip()
            if not email or email.lower() in ("nan", "none", ""):
                continue
        if not _EMAIL_RE.fullmatch(email):
            continue

        # Build name
        name = ""
        if first_col or last_col:
            first = str(row.get(first_col, "") or "").strip() if first_col else ""
            last = str(row.get(last_col, "") or "").strip() if last_col else ""
            parts = [p for p in [first, last] if p and p.lower() != "nan"]
            name = " ".join(parts)
        if not name and name_col:
            name = str(row.get(name_col, "") or "").strip()
        if not name or name.lower() == "nan":
            name = email.split("@")[0]

        branch = str(row.get(branch_col, "") or "").strip() if branch_col else ""
        sales_rep = str(row.get(salesrep_col, "") or "").strip() if salesrep_col else ""
        dept = str(row.get(dept_col, "") or "").strip() if dept_col else ""
        company = str(row.get(company_col, "") or "").strip() if company_col else ""
        title = str(row.get(title_col, "") or "").strip() if title_col else ""

        raw = {str(k): str(v) for k, v in row.items()
               if v is not None and str(v).strip() not in ("", "nan", "None")}

        results.append({
            "name": name,
            "email": email.lower(),
            "branch": branch if branch.lower() != "nan" else "",
            "sales_rep": sales_rep if sales_rep.lower() != "nan" else "",
            "department": dept if dept.lower() != "nan" else "",
            "company": company if company.lower() != "nan" else "",
            "title": title if title.lower() != "nan" else "",
            "source_file": source_file,
            "raw_row": json.dumps(raw, ensure_ascii=False),
        })
    return results


def extract_from_csv_bytes(file_bytes: bytes, filename: str = "") -> list[dict]:
    try:
        df = pd.read_csv(io.BytesIO(file_bytes), dtype=str, on_bad_lines="skip")
        return _extract_from_dataframe(df, filename)
    except Exception as e:
        log.error("CSV contact extraction error: %s", e)
        return []


def extract_from_excel_bytes(file_bytes: bytes, filename: str = "") -> list[dict]:
    results = []
    try:
        sheets = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None, dtype=str)
        for sheet_name, df in sheets.items():
            results.extend(_extract_from_dataframe(df, filename))
    except Exception as e:
        log.error("Excel contact extraction error: %s", e)
    return results


def extract_from_docx_bytes(file_bytes: bytes, filename: str = "") -> list[dict]:
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        text = "\n".join(p.text for p in doc.paragraphs)
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + "\t".join(cell.text for cell in row.cells)
        return _extract_from_text(text, filename)
    except ImportError:
        log.warning("python-docx not installed; skipping DOCX extraction")
        return []
    except Exception as e:
        log.error("DOCX contact extraction error: %s", e)
        return []


def extract_from_pdf_bytes(file_bytes: bytes, filename: str = "") -> list[dict]:
    try:
        import pdfplumber
        text = ""
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                text += (page.extract_text() or "") + "\n"
        return _extract_from_text(text, filename)
    except ImportError:
        log.warning("pdfplumber not installed; skipping PDF extraction")
        return []
    except Exception as e:
        log.error("PDF contact extraction error: %s", e)
        return []


def extract_from_txt_bytes(file_bytes: bytes, filename: str = "") -> list[dict]:
    try:
        text = file_bytes.decode("utf-8", errors="replace")
        return _extract_from_text(text, filename)
    except Exception as e:
        log.error("TXT contact extraction error: %s", e)
        return []


def _extract_from_text(text: str, filename: str = "") -> list[dict]:
    """Find email addresses in text and associate nearby context as name."""
    results = []
    lines = text.splitlines()
    for i, line in enumerate(lines):
        emails = _EMAIL_RE.findall(line)
        for email in emails:
            # Look for a name in the same line or adjacent lines
            ctx_lines = lines[max(0, i - 1):i + 2]
            ctx = " ".join(ctx_lines)
            # Remove the email itself from context
            ctx = ctx.replace(email, "").strip()
            # Try to find name-like tokens (no digits, length > 2)
            tokens = [t for t in re.split(r"[\s,;:|]+", ctx)
                      if len(t) > 2 and not re.search(r"\d", t)]
            name = " ".join(tokens[:3]) if tokens else email.split("@")[0]
            results.append({
                "name": name,
                "email": email.lower(),
                "branch": "",
                "sales_rep": "",
                "department": "",
                "company": "",
                "title": "",
                "source_file": filename,
                "raw_row": json.dumps({"line": line.strip()}),
            })
    # Deduplicate by email
    seen = set()
    deduped = []
    for r in results:
        if r["email"] not in seen:
            seen.add(r["email"])
            deduped.append(r)
    return deduped


def extract_contacts(file_bytes: bytes, filename: str) -> list[dict]:
    """Dispatch to the right extractor based on file extension."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".csv":
        return extract_from_csv_bytes(file_bytes, filename)
    elif ext in (".xlsx", ".xls"):
        return extract_from_excel_bytes(file_bytes, filename)
    elif ext == ".docx":
        return extract_from_docx_bytes(file_bytes, filename)
    elif ext == ".pdf":
        return extract_from_pdf_bytes(file_bytes, filename)
    elif ext == ".txt":
        return extract_from_txt_bytes(file_bytes, filename)
    else:
        log.warning("Unsupported contact file type: %s", ext)
        return []
